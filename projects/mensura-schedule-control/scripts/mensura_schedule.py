#!/usr/bin/env python3
"""mensura-schedule — CLI interna do Mensura Schedule Control.

Agent-native command layer for SharePoint schedule discovery, Supabase ingestion,
and critical-path analytics.
"""
from __future__ import annotations

import argparse
import base64
import json
import hashlib
import os
import re
import subprocess
import sys
from datetime import datetime, date
from pathlib import Path
from typing import Any

import psycopg2
import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKSPACE = Path("/root/.openclaw/workspace")
ENV_FILE = WORKSPACE / "memory/context/mensura_schedule_supabase.env"
MENSURA_GRAPH_CONFIG = WORKSPACE / "config/ms-graph.json"
GRAPH = "https://graph.microsoft.com/v1.0"

# Projetos excluídos do relatório executivo por decisão operacional do Alê.
# MELICITA é duplicado de MELICITA_R1; manter MELICITA_R1 como versão válida.
EXECUTIVE_EXCLUDED_PROJECT_CODES = (
    "DOPPIO",
    "MELICITA",
    "ELEV_ALTO_DO_IPIRANGA",
    "SOFITEL_DIRETOR",
    "CCN_BIOMA",
    "DF345_DIOGO_DE_FARIA",
)
EXECUTIVE_EXCLUDED_SQL = "('DOPPIO','MELICITA','ELEV_ALTO_DO_IPIRANGA','SOFITEL_DIRETOR','CCN_BIOMA','DF345_DIOGO_DE_FARIA')"


def norm(s: Any) -> str:
    s = "" if s is None else str(s).strip().lower()
    s = s.replace("ì", "í")
    s = re.sub(r"\s+", " ", s)
    return s

def load_env_file(path: Path = ENV_FILE) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k, v)


def db_conn():
    load_env_file()
    dsn = os.environ.get("MENSURA_SCHEDULE_DATABASE_URL") or os.environ.get("DATABASE_URL")
    if dsn:
        return psycopg2.connect(dsn)
    password = os.environ.get("MENSURA_SCHEDULE_DB_PASSWORD")
    if not password:
        raise SystemExit(f"Missing MENSURA_SCHEDULE_DB_PASSWORD; expected local env at {ENV_FILE}")
    host = os.environ.get("MENSURA_SCHEDULE_DB_HOST", "db.ckmuyvbacgdidmiccvif.supabase.co")
    return psycopg2.connect(host=host, port=5432, dbname="postgres", user="postgres", password=password, sslmode="require")


def print_rows(rows: list[tuple], headers: list[str], as_json: bool = False) -> None:
    if as_json:
        print(json.dumps([dict(zip(headers, r)) for r in rows], ensure_ascii=False, indent=2, default=str))
        return
    print("\t".join(headers))
    for r in rows:
        print("\t".join("" if v is None else str(v) for v in r))



def parse_project_date(v: Any):
    if v in (None, ""):
        return None
    if hasattr(v, "isoformat") and v.__class__.__name__ in {"date", "datetime"}:
        return v.date() if hasattr(v, "date") else v
    txt = str(v).strip()
    if not txt or norm(txt) in {"nd", "n/d", "na", "n.a.", "-"}:
        return None
    txt = re.sub(r"^(seg|ter|qua|qui|sex|sáb|sab|dom)\.?\s+", "", txt, flags=re.I).strip()
    m = re.search(r"\d{1,2}/\d{1,2}/\d{2,4}", txt)
    if m:
        txt = m.group(0)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(txt[:19], fmt).date()
        except Exception:
            pass
    return None


def raw_get_alias(raw: dict[str, Any], aliases: list[str]):
    nmap = {norm(k): k for k in raw.keys()}
    for alias in aliases:
        k = nmap.get(norm(alias))
        if k is not None:
            return raw.get(k)
    return None

def cmd_db_counts(args):
    sql = """
    select 'projects' object, count(*) from schedule.projects
    union all select 'import_jobs', count(*) from raw.import_jobs
    union all select 'source_rows', count(*) from raw.source_rows
    union all select 'schedule_versions', count(*) from schedule.schedule_versions
    union all select 'activity_identities', count(*) from schedule.activity_identities
    union all select 'activity_versions', count(*) from schedule.activity_versions
    union all select 'critical_current', count(*) from analytics.v_current_critical_activities
    order by 1;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql)
        print_rows(cur.fetchall(), ["object", "count"], args.json)


def cmd_critical_summary(args):
    sql = """
    select
      project_code,
      activities_total,
      critical_total,
      critical_open,
      critical_overdue,
      round(avg_percent_complete::numeric, 2) as avg_percent_complete,
      projected_delay_days_vs_baseline
    from analytics.v_project_schedule_metrics_current
    order by critical_total desc nulls last, project_code
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.limit,))
        print_rows(cur.fetchall(), ["project", "activities", "critical", "critical_open", "critical_overdue", "avg_pct", "delay_vs_baseline"], args.json)


def cmd_quality_report(args):
    sql = """
    with latest as (
      select distinct on (sv.project_id)
        sv.id, sv.project_id, p.code project_code, sv.version_label, sv.data_date
      from schedule.schedule_versions sv
      join schedule.projects p on p.id = sv.project_id
      order by sv.project_id, sv.data_date desc nulls last, sv.created_at desc
    ), q as (
      select
        l.project_code,
        l.version_label,
        l.data_date,
        count(*) activities,
        count(*) filter (where av.current_start is null) missing_current_start,
        count(*) filter (where av.current_finish is null) missing_current_finish,
        count(*) filter (where av.percent_complete is null) missing_percent,
        count(*) filter (where av.is_critical is null) missing_critical_flag,
        count(*) filter (where av.baseline_finish is null) missing_baseline_finish,
        count(*) filter (where av.current_finish < av.current_start) inverted_current_dates,
        count(*) filter (where av.percent_complete < 0 or av.percent_complete > 100) invalid_percent
      from latest l
      join schedule.activity_versions av on av.schedule_version_id = l.id
      group by l.project_code, l.version_label, l.data_date
    )
    select *,
      round(100 * (1 - (missing_current_finish + missing_percent + missing_critical_flag + inverted_current_dates + invalid_percent)::numeric / nullif(activities * 5, 0)), 2) quality_score
    from q
    order by quality_score asc nulls last, activities desc
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.limit,))
        print_rows(cur.fetchall(), ["project", "version", "data_date", "activities", "missing_start", "missing_finish", "missing_pct", "missing_critical", "missing_baseline", "inverted_dates", "invalid_pct", "quality_score"], args.json)


def cmd_baseline_compare(args):
    sql = """
    with latest as (
      select distinct on (sv.project_id)
        sv.id, sv.project_id, p.code project_code, sv.version_label, sv.data_date
      from schedule.schedule_versions sv
      join schedule.projects p on p.id = sv.project_id
      order by sv.project_id, sv.data_date desc nulls last, sv.created_at desc
    )
    select
      l.project_code,
      count(*) filter (where av.baseline_finish is not null and av.current_finish is not null) comparable_activities,
      count(*) filter (where av.baseline_finish is not null and av.current_finish > av.baseline_finish) delayed_activities,
      count(*) filter (where av.is_critical and av.baseline_finish is not null and av.current_finish > av.baseline_finish) delayed_critical_activities,
      max(av.current_finish - av.baseline_finish) max_finish_slip_days,
      round(avg((av.current_finish - av.baseline_finish))::numeric, 2) avg_finish_slip_days
    from latest l
    join schedule.activity_versions av on av.schedule_version_id = l.id
    where (%s is null or l.project_code = %s)
    group by l.project_code
    order by delayed_critical_activities desc nulls last, delayed_activities desc nulls last
    limit %s;
    """
    project = args.project
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (project, project, args.limit))
        print_rows(cur.fetchall(), ["project", "comparable", "delayed", "delayed_critical", "max_slip_days", "avg_slip_days"], args.json)


def cmd_weekly_metrics_preview(args):
    sql = """
    select
      project_code,
      activities_total,
      round(avg_percent_complete::numeric, 2) avg_percent_complete,
      critical_total,
      critical_open,
      critical_overdue,
      max_current_finish as projected_finish_current,
      projected_delay_days_vs_baseline
    from analytics.v_project_schedule_metrics_current
    order by critical_overdue desc nulls last, projected_delay_days_vs_baseline desc nulls last, project_code
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.limit,))
        print_rows(cur.fetchall(), ["project", "activities", "avg_pct", "critical", "critical_open", "critical_overdue", "projected_finish", "delay_vs_baseline"], args.json)


def cmd_audit_orphans(args):
    sql = """
    select
      ij.status,
      count(*) jobs,
      min(ij.imported_at) first_imported_at,
      max(ij.imported_at) last_imported_at
    from raw.import_jobs ij
    left join schedule.schedule_versions sv on sv.import_job_id = ij.id
    where sv.id is null
    group by ij.status
    order by jobs desc;
    """
    detail_sql = """
    select ij.id, ij.status, ij.source_filename, ij.imported_at, ij.data_date
    from raw.import_jobs ij
    left join schedule.schedule_versions sv on sv.import_job_id = ij.id
    where sv.id is null
    order by ij.imported_at desc
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        if args.detail:
            cur.execute(detail_sql, (args.limit,))
            print_rows(cur.fetchall(), ["id", "status", "source_filename", "imported_at", "data_date"], args.json)
        else:
            cur.execute(sql)
            print_rows(cur.fetchall(), ["status", "jobs", "first_imported_at", "last_imported_at"], args.json)



def cmd_project_classify(args):
    sql = """
    with latest as (
      select distinct on (sv.project_id)
        sv.id schedule_version_id, sv.project_id, p.code project_code, p.name project_name,
        sv.version_label, sv.data_date, sv.created_at
      from schedule.schedule_versions sv
      join schedule.projects p on p.id = sv.project_id
      order by sv.project_id, sv.data_date desc nulls last, sv.created_at desc
    ), metrics as (
      select
        l.project_id,
        count(av.id) activities,
        count(*) filter (where av.is_critical is true) critical,
        round(avg(av.percent_complete) filter (where av.percent_complete is not null)::numeric, 2) avg_pct,
        min(av.current_start) min_start,
        max(av.current_finish) max_finish,
        count(*) filter (where av.current_finish is null) missing_finish,
        count(*) filter (where av.is_critical is null) missing_critical
      from latest l
      join schedule.activity_versions av on av.schedule_version_id = l.schedule_version_id
      group by l.project_id
    ), dupkey as (
      select
        coalesce(m.activities,0) activities,
        coalesce(m.critical,0) critical,
        m.min_start,
        m.max_finish,
        count(*) projects_same_shape
      from latest l
      join metrics m on m.project_id = l.project_id
      group by coalesce(m.activities,0), coalesce(m.critical,0), m.min_start, m.max_finish
    )
    select
      l.project_code,
      l.project_name,
      l.version_label,
      l.data_date,
      m.activities,
      m.critical,
      m.avg_pct,
      case
        when l.project_code ~* '(^|_)TESTE($|_)|OBRA_MODELO|MODELO|TEMPLATE|SCHDULE|ARQUIVOS_AUXILIARES' then 'exclude_model_or_auxiliary'
        when d.projects_same_shape > 1 and l.project_code ~* 'TESTE|COPY|C[ÓO]PIA|ARQUIVOS_AUXILIARES' then 'exclude_probable_duplicate'
        when m.activities < 20 then 'review_too_few_activities'
        when (m.missing_finish::numeric / nullif(m.activities,0)) > 0.25 then 'review_low_date_quality'
        when (m.missing_critical::numeric / nullif(m.activities,0)) > 0.25 then 'review_low_critical_quality'
        else 'valid_for_analytics'
      end as classification,
      d.projects_same_shape,
      m.missing_finish,
      m.missing_critical
    from latest l
    join metrics m on m.project_id = l.project_id
    join dupkey d on d.activities = coalesce(m.activities,0)
      and d.critical = coalesce(m.critical,0)
      and d.min_start is not distinct from m.min_start
      and d.max_finish is not distinct from m.max_finish
    order by classification desc, m.critical desc nulls last, l.project_code
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.limit,))
        print_rows(cur.fetchall(), ["project", "name", "version", "data_date", "activities", "critical", "avg_pct", "classification", "same_shape", "missing_finish", "missing_critical"], args.json)


def cmd_baseline_discovery(args):
    sql = """
    with by_version as (
      select
        p.code project_code,
        sv.version_label,
        sv.data_date,
        sv.is_baseline_candidate,
        count(av.id) activities,
        count(*) filter (where av.baseline_finish is not null) baseline_finish_count,
        count(*) filter (where av.current_finish is not null) current_finish_count,
        max(av.baseline_finish) max_baseline_finish,
        max(av.current_finish) max_current_finish,
        case
          when sv.is_baseline_candidate then 'explicit_candidate'
          when sv.version_label ~* 'baseline|base.?line|linha.?base|lb|r0|rev.?0|original' then 'label_candidate'
          when count(*) filter (where av.baseline_finish is not null) > 0 then 'activity_baseline_fields'
          else 'no_baseline_signal'
        end baseline_signal
      from schedule.schedule_versions sv
      join schedule.projects p on p.id = sv.project_id
      join schedule.activity_versions av on av.schedule_version_id = sv.id
      group by p.code, sv.id
    )
    select
      project_code,
      version_label,
      data_date,
      baseline_signal,
      activities,
      baseline_finish_count,
      round(100 * q.baseline_finish_count::numeric / nullif(q.activities,0), 2) baseline_finish_coverage_pct,
      max_baseline_finish,
      max_current_finish,
      (max_current_finish - max_baseline_finish) project_finish_variance_days
    from by_version
    where (%s is false or baseline_signal <> 'no_baseline_signal')
    order by project_code, case baseline_signal when 'explicit_candidate' then 1 when 'label_candidate' then 2 when 'activity_baseline_fields' then 3 else 4 end, data_date
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.only_candidates, args.limit))
        print_rows(cur.fetchall(), ["project", "version", "data_date", "baseline_signal", "activities", "baseline_finish", "coverage_pct", "max_baseline_finish", "max_current_finish", "finish_variance_days"], args.json)


def cmd_risk_report(args):
    sql = """
    with q as (
      with latest as (
        select distinct on (sv.project_id) sv.id, sv.project_id, p.code project_code, sv.version_label, sv.data_date
        from schedule.schedule_versions sv
        join schedule.projects p on p.id = sv.project_id
        order by sv.project_id, sv.data_date desc nulls last, sv.created_at desc
      )
      select
        l.project_id,
        count(*) activities,
        count(*) filter (where av.current_finish is null) missing_finish,
        count(*) filter (where av.percent_complete is null) missing_percent,
        count(*) filter (where av.is_critical is null) missing_critical,
        count(*) filter (where av.current_finish < av.current_start) inverted_dates,
        count(*) filter (where av.percent_complete < 0 or av.percent_complete > 100) invalid_percent
      from latest l
      join schedule.activity_versions av on av.schedule_version_id = l.id
      group by l.project_id
    )
    select
      m.project_code,
      m.version_label,
      m.data_date,
      m.activities_total,
      m.critical_total,
      m.critical_open,
      m.critical_overdue,
      m.activities_overdue,
      round(m.avg_percent_complete::numeric, 2) avg_pct,
      m.max_current_finish,
      m.projected_delay_days_vs_baseline,
      round(100 * (1 - (q.missing_finish + q.missing_percent + q.missing_critical + q.inverted_dates + q.invalid_percent)::numeric / nullif(q.activities * 5, 0)), 2) quality_score,
      (
        coalesce(m.critical_overdue,0) * 5
        + coalesce(m.activities_overdue,0) * 2
        + coalesce(m.critical_open,0) * 0.5
        + greatest(coalesce(m.projected_delay_days_vs_baseline,0),0) * 1.5
        + greatest(85 - coalesce(round(100 * (1 - (q.missing_finish + q.missing_percent + q.missing_critical + q.inverted_dates + q.invalid_percent)::numeric / nullif(q.activities * 5, 0)), 2),85),0)
      )::numeric(12,2) risk_score,
      case
        when coalesce(m.critical_overdue,0) > 0 or coalesce(m.projected_delay_days_vs_baseline,0) > 30 then 'CRITICO'
        when coalesce(m.activities_overdue,0) > 0 or coalesce(m.critical_open,0) > 100 then 'ATENCAO'
        when round(100 * (1 - (q.missing_finish + q.missing_percent + q.missing_critical + q.inverted_dates + q.invalid_percent)::numeric / nullif(q.activities * 5, 0)), 2) < 80 then 'DADO_FRACO'
        else 'OK'
      end risk_level
    from analytics.v_project_schedule_metrics_current m
    join q on q.project_id = m.project_id
    order by risk_score desc nulls last, m.critical_total desc nulls last, m.project_code
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.limit,))
        print_rows(cur.fetchall(), ["project", "version", "data_date", "activities", "critical", "critical_open", "critical_overdue", "activities_overdue", "avg_pct", "projected_finish", "delay_vs_baseline", "quality_score", "risk_score", "risk_level"], args.json)


def cmd_populate_weekly_metrics(args):
    sql = """
    insert into analytics.project_weekly_metrics (
      project_id, week_start, data_date, planned_percent, actual_percent,
      critical_activities_delayed_count, negative_float_activities_count,
      forecast_finish_date_current, delay_days_vs_contract, metrics
    )
    select
      m.project_id,
      date_trunc('week', m.data_date)::date as week_start,
      m.data_date,
      null::numeric as planned_percent,
      round(m.avg_percent_complete::numeric, 2) as actual_percent,
      m.critical_overdue::integer as critical_activities_delayed_count,
      0::integer as negative_float_activities_count,
      m.max_current_finish as forecast_finish_date_current,
      case when p.contract_finish_date is not null and m.max_current_finish is not null
        then (m.max_current_finish - p.contract_finish_date)::numeric
        else m.projected_delay_days_vs_baseline::numeric
      end as delay_days_vs_contract,
      jsonb_build_object(
        'source', 'mensura-schedule CLI populate-weekly-metrics',
        'version_label', m.version_label,
        'activities_total', m.activities_total,
        'critical_total', m.critical_total,
        'critical_open', m.critical_open,
        'activities_overdue', m.activities_overdue,
        'avg_finish_variance_days', m.avg_finish_variance_days,
        'projected_delay_days_vs_baseline', m.projected_delay_days_vs_baseline
      ) as metrics
    from analytics.v_project_schedule_metrics_current m
    join schedule.projects p on p.id = m.project_id
    on conflict (project_id, week_start) do update set
      data_date = excluded.data_date,
      planned_percent = excluded.planned_percent,
      actual_percent = excluded.actual_percent,
      critical_activities_delayed_count = excluded.critical_activities_delayed_count,
      negative_float_activities_count = excluded.negative_float_activities_count,
      forecast_finish_date_current = excluded.forecast_finish_date_current,
      delay_days_vs_contract = excluded.delay_days_vs_contract,
      metrics = excluded.metrics;
    """
    preview_sql = """
    select
      m.project_code,
      date_trunc('week', m.data_date)::date week_start,
      m.data_date,
      round(m.avg_percent_complete::numeric, 2) actual_percent,
      m.critical_overdue,
      m.max_current_finish forecast_finish_date_current,
      m.projected_delay_days_vs_baseline
    from analytics.v_project_schedule_metrics_current m
    order by m.project_code
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        if args.execute:
            cur.execute(sql)
            affected = cur.rowcount
            conn.commit()
            if args.json:
                print(json.dumps({"mode": "execute", "upserted_or_updated": affected}, ensure_ascii=False, indent=2))
            else:
                print(f"upserted_or_updated={affected}")
        cur.execute(preview_sql, (args.limit,))
        print_rows(cur.fetchall(), ["project", "week_start", "data_date", "actual_pct", "critical_overdue", "forecast_finish", "delay_vs_baseline"], args.json)


def cmd_data_readiness(args):
    sql = """
    with latest as (
      select distinct on (sv.project_id)
        sv.id schedule_version_id, sv.project_id, p.code project_code, p.name project_name,
        sv.version_label, sv.data_date, sv.created_at
      from schedule.schedule_versions sv
      join schedule.projects p on p.id = sv.project_id
      order by sv.project_id, sv.data_date desc nulls last, sv.created_at desc
    ), q as (
      select
        l.project_id,
        l.project_code,
        l.project_name,
        l.version_label,
        l.data_date,
        count(av.id) activities,
        count(*) filter (where av.current_start is not null) current_start_count,
        count(*) filter (where av.current_finish is not null) current_finish_count,
        count(*) filter (where av.percent_complete is not null) percent_count,
        count(*) filter (where av.is_critical is not null) critical_flag_count,
        count(*) filter (where av.baseline_finish is not null) baseline_finish_count,
        count(*) filter (where av.is_critical is true) critical_total,
        count(*) filter (where av.is_critical is true and av.actual_finish is null and coalesce(av.percent_complete,0) < 100) critical_open,
        count(*) filter (where av.current_finish < av.current_start) inverted_dates,
        count(*) filter (where av.percent_complete < 0 or av.percent_complete > 100) invalid_percent,
        max(av.current_finish) max_current_finish,
        max(av.baseline_finish) max_baseline_finish
      from latest l
      join schedule.activity_versions av on av.schedule_version_id = l.schedule_version_id
      group by l.project_id, l.project_code, l.project_name, l.version_label, l.data_date
    ), shape as (
      select activities, critical_total, max_current_finish, count(*) same_shape_projects
      from q
      group by activities, critical_total, max_current_finish
    ), scored as (
      select
        q.*,
        round(100 * q.current_start_count::numeric / nullif(q.activities,0), 2) current_start_coverage,
        round(100 * q.current_finish_count::numeric / nullif(q.activities,0), 2) current_finish_coverage,
        round(100 * q.percent_count::numeric / nullif(q.activities,0), 2) percent_coverage,
        round(100 * q.critical_flag_count::numeric / nullif(q.activities,0), 2) critical_flag_coverage,
        round(100 * q.baseline_finish_count::numeric / nullif(q.activities,0), 2) baseline_finish_coverage,
        s.same_shape_projects,
        case
          when q.project_code ~* '(^|_)TESTE($|_)|OBRA_MODELO|MODELO|TEMPLATE|SCHDULE|ARQUIVOS_AUXILIARES' then true
          else false
        end is_model_or_auxiliary,
        (
          least(round(100 * q.current_finish_count::numeric / nullif(q.activities,0), 2),100) * 0.30
          + least(round(100 * q.percent_count::numeric / nullif(q.activities,0), 2),100) * 0.20
          + least(round(100 * q.critical_flag_count::numeric / nullif(q.activities,0), 2),100) * 0.20
          + least(round(100 * q.baseline_finish_count::numeric / nullif(q.activities,0), 2),100) * 0.20
          + case when q.inverted_dates = 0 and q.invalid_percent = 0 then 10 else 0 end
        )::numeric(8,2) readiness_score
      from q
      join shape s on s.activities = q.activities
        and s.critical_total = q.critical_total
        and s.max_current_finish is not distinct from q.max_current_finish
    )
    select
      project_code,
      version_label,
      data_date,
      activities,
      critical_total,
      critical_open,
      current_finish_coverage,
      percent_coverage,
      critical_flag_coverage,
      baseline_finish_coverage,
      inverted_dates,
      invalid_percent,
      same_shape_projects,
      readiness_score,
      case
        when is_model_or_auxiliary then 'EXCLUIR_MODELO_AUXILIAR'
        when readiness_score >= 85 and baseline_finish_coverage >= 80 then 'PRONTO_PREDITIVO'
        when readiness_score >= 75 and baseline_finish_coverage >= 20 then 'PRONTO_ANALISE_CONTROLE'
        when current_finish_coverage < 80 then 'BLOQUEADO_DATAS_ATUAIS'
        when baseline_finish_coverage < 20 then 'BLOQUEADO_BASELINE'
        when critical_flag_coverage < 80 then 'BLOQUEADO_CRITICA'
        else 'REVISAO_MANUAL'
      end readiness_level,
      case
        when is_model_or_auxiliary then 'Excluir do universo analítico ou mapear manualmente se for obra real.'
        when current_finish_coverage < 80 then 'Revisar mapeamento/importação das datas atuais de término no Excel/MS Project.'
        when baseline_finish_coverage < 20 then 'Selecionar/importar baseline aprovada e registrar em schedule.baselines.'
        when critical_flag_coverage < 80 then 'Revisar captura da coluna Crítica do MS Project.'
        when readiness_score >= 85 and baseline_finish_coverage >= 80 then 'Pode entrar na camada preditiva inicial.'
        when readiness_score >= 75 then 'Pode entrar em análise de controle; predição depende de baseline/histórico.'
        else 'Revisão manual de qualidade e identidade das atividades.'
      end recommended_action
    from scored
    order by
      case
        when is_model_or_auxiliary then 5
        when readiness_score >= 85 and baseline_finish_coverage >= 80 then 1
        when readiness_score >= 75 and baseline_finish_coverage >= 20 then 2
        when current_finish_coverage < 80 then 3
        else 4
      end,
      readiness_score desc nulls last,
      critical_total desc nulls last,
      project_code
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.limit,))
        print_rows(cur.fetchall(), [
            "project", "version", "data_date", "activities", "critical", "critical_open",
            "current_finish_cov", "percent_cov", "critical_cov", "baseline_cov",
            "inverted_dates", "invalid_pct", "same_shape", "readiness_score", "readiness_level", "recommended_action"
        ], args.json)


def cmd_repair_date_fields(args):
    aliases = {
        "baseline_start": ["início da linha de base", "inicio da linha de base", "início da linha de base1", "início da linha de base (meta)"],
        "baseline_finish": ["término da linha de base", "termino da linha de base", "término da linha de base1", "término da linha de base (meta)"],
        "current_start": ["início", "inicio"],
        "current_finish": ["término", "termino"],
        "actual_start": ["início real", "inicio real"],
        "actual_finish": ["término real", "termino real"],
    }
    select_sql = """
      select av.id, av.raw_payload,
             av.baseline_start, av.baseline_finish, av.current_start, av.current_finish, av.actual_start, av.actual_finish
      from schedule.activity_versions av
      where av.raw_payload is not null
        and (%s is true or av.current_finish is null or av.current_start is null or av.baseline_finish is null or av.baseline_start is null)
      limit %s;
    """
    update_sql = """
      update schedule.activity_versions
      set baseline_start=%s, baseline_finish=%s, current_start=%s, current_finish=%s, actual_start=%s, actual_finish=%s
      where id=%s;
    """
    changed = 0
    scanned = 0
    examples = []
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(select_sql, (args.overwrite, args.limit))
        rows = cur.fetchall()
        for row in rows:
            scanned += 1
            av_id, raw, bs, bf, cs, cf, astart, afinish = row
            raw = raw or {}
            parsed = {
                "baseline_start": parse_project_date(raw_get_alias(raw, aliases["baseline_start"])),
                "baseline_finish": parse_project_date(raw_get_alias(raw, aliases["baseline_finish"])),
                "current_start": parse_project_date(raw_get_alias(raw, aliases["current_start"])),
                "current_finish": parse_project_date(raw_get_alias(raw, aliases["current_finish"])),
                "actual_start": parse_project_date(raw_get_alias(raw, aliases["actual_start"])),
                "actual_finish": parse_project_date(raw_get_alias(raw, aliases["actual_finish"])),
            }
            new_vals = {
                "baseline_start": parsed["baseline_start"] if (args.overwrite or bs is None) else bs,
                "baseline_finish": parsed["baseline_finish"] if (args.overwrite or bf is None) else bf,
                "current_start": parsed["current_start"] if (args.overwrite or cs is None) else cs,
                "current_finish": parsed["current_finish"] if (args.overwrite or cf is None) else cf,
                "actual_start": parsed["actual_start"] if (args.overwrite or astart is None) else astart,
                "actual_finish": parsed["actual_finish"] if (args.overwrite or afinish is None) else afinish,
            }
            old_vals = {"baseline_start": bs, "baseline_finish": bf, "current_start": cs, "current_finish": cf, "actual_start": astart, "actual_finish": afinish}
            if new_vals != old_vals:
                changed += 1
                if len(examples) < 5:
                    examples.append({"id": str(av_id), "old": old_vals, "new": new_vals})
                if args.execute:
                    cur.execute(update_sql, (new_vals["baseline_start"], new_vals["baseline_finish"], new_vals["current_start"], new_vals["current_finish"], new_vals["actual_start"], new_vals["actual_finish"], av_id))
        if args.execute:
            conn.commit()
    payload = {"mode": "execute" if args.execute else "dry_run", "scanned": scanned, "changed": changed, "examples": examples}
    if args.json:
        print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))
    else:
        print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))


def cmd_analytics_universe(args):
    sql = """
    with latest as (
      select distinct on (sv.project_id)
        sv.id schedule_version_id, sv.project_id, p.code project_code, p.name project_name,
        sv.version_label, sv.data_date
      from schedule.schedule_versions sv
      join schedule.projects p on p.id = sv.project_id
      order by sv.project_id, sv.data_date desc nulls last, sv.created_at desc
    ), q as (
      select
        l.project_id, l.project_code, l.project_name, l.version_label, l.data_date,
        count(av.id) activities,
        count(*) filter (where av.current_finish is not null) current_finish_count,
        count(*) filter (where av.percent_complete is not null) percent_count,
        count(*) filter (where av.is_critical is not null) critical_flag_count,
        count(*) filter (where av.baseline_finish is not null) baseline_finish_count,
        count(*) filter (where av.is_critical is true) critical_total,
        max(av.current_finish) max_current_finish,
        max(av.baseline_finish) max_baseline_finish
      from latest l
      join schedule.activity_versions av on av.schedule_version_id = l.schedule_version_id
      group by l.project_id, l.project_code, l.project_name, l.version_label, l.data_date
    ), scored as (
      select *,
        round(100 * current_finish_count::numeric / nullif(activities,0), 2) current_finish_coverage,
        round(100 * percent_count::numeric / nullif(activities,0), 2) percent_coverage,
        round(100 * critical_flag_count::numeric / nullif(activities,0), 2) critical_flag_coverage,
        round(100 * baseline_finish_count::numeric / nullif(activities,0), 2) baseline_finish_coverage,
        case when project_code ~* '(^|_)TESTE($|_)|OBRA_MODELO|MODELO|TEMPLATE|SCHDULE|ARQUIVOS_AUXILIARES' then true else false end is_model_or_auxiliary
      from q
    )
    select
      project_code,
      project_name,
      version_label,
      data_date,
      activities,
      critical_total,
      current_finish_coverage,
      baseline_finish_coverage,
      (max_current_finish - max_baseline_finish) delay_vs_baseline,
      case
        when is_model_or_auxiliary then 'exclude'
        when exists (select 1 from portfolio.project_registry pr where pr.company='Mensura' and pr.schedule_project_id=project_id and pr.include_in_forecast is false and pr.operational_status in ('excluded','duplicate')) then 'exclude'
        when current_finish_coverage >= 95 and baseline_finish_coverage >= 80 then 'predictive'
        when current_finish_coverage >= 95 then 'control_only'
        else 'blocked'
      end as universe_status,
      case
        when is_model_or_auxiliary then 'modelo/auxiliar/teste; excluir do universo analítico'
        when exists (select 1 from portfolio.project_registry pr where pr.company='Mensura' and pr.schedule_project_id=project_id and pr.include_in_forecast is false and pr.operational_status in ('excluded','duplicate')) then 'excluído/duplicado conforme portfolio.project_registry'
        when current_finish_coverage >= 95 and baseline_finish_coverage >= 80 then 'apto para risk-report, baseline compare e forecast inicial'
        when current_finish_coverage >= 95 then 'apto para controle; baseline insuficiente para previsão robusta'
        else 'bloqueado por qualidade de datas'
      end as note
    from scored
    where (%s is false or (not is_model_or_auxiliary and exists (select 1 from portfolio.project_registry pr where pr.company='Mensura' and pr.schedule_project_id=project_id and pr.include_in_forecast is true)))
    order by case
        when is_model_or_auxiliary then 4
        when current_finish_coverage >= 95 and baseline_finish_coverage >= 80 then 1
        when current_finish_coverage >= 95 then 2
        else 3 end,
      coalesce((max_current_finish - max_baseline_finish),0) desc, project_code
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.analytics_only, args.limit))
        print_rows(cur.fetchall(), ["project", "name", "version", "data_date", "activities", "critical", "current_finish_cov", "baseline_cov", "delay_vs_baseline", "universe_status", "note"], args.json)


def cmd_forecast_initial(args):
    base_sql = """
    with m as (
      select m.*, p.contract_finish_date,
        case when p.code ~* '(^|_)TESTE($|_)|OBRA_MODELO|MODELO|TEMPLATE|SCHDULE|ARQUIVOS_AUXILIARES' then true else false end is_model_or_auxiliary
      from analytics.v_project_schedule_metrics_current m
      join schedule.projects p on p.id = m.project_id
    ), f as (
      select
        project_id,
        project_code,
        data_date,
        max_current_finish predicted_finish_date,
        greatest(coalesce(projected_delay_days_vs_baseline, 0), 0)::numeric expected_delay_days,
        case
          when projected_delay_days_vs_baseline is null then null
          when projected_delay_days_vs_baseline <= 0 then 0.15
          when projected_delay_days_vs_baseline <= 15 then 0.45
          when projected_delay_days_vs_baseline <= 30 then 0.65
          when projected_delay_days_vs_baseline <= 60 then 0.80
          else 0.90
        end::numeric p_finish_after_contract,
        jsonb_build_array(
          jsonb_build_object('driver','projected_delay_days_vs_baseline','value',projected_delay_days_vs_baseline),
          jsonb_build_object('driver','critical_open','value',critical_open),
          jsonb_build_object('driver','critical_overdue','value',critical_overdue),
          jsonb_build_object('driver','avg_percent_complete','value',round(avg_percent_complete::numeric,2))
        ) drivers,
        jsonb_build_object(
          'version_label', version_label,
          'activities_total', activities_total,
          'critical_total', critical_total,
          'critical_open', critical_open,
          'critical_overdue', critical_overdue,
          'activities_overdue', activities_overdue,
          'avg_percent_complete', round(avg_percent_complete::numeric,2),
          'max_current_finish', max_current_finish,
          'max_baseline_finish', max_baseline_finish,
          'projected_delay_days_vs_baseline', projected_delay_days_vs_baseline
        ) feature_snapshot
      from m
      where not is_model_or_auxiliary
        and max_current_finish is not null
        and max_baseline_finish is not null
        and max_current_finish >= current_date
        and exists (select 1 from portfolio.project_registry pr where pr.company='Mensura' and pr.schedule_project_id=m.project_id and pr.include_in_forecast is true)
    )
    """
    preview_sql = base_sql + """
    select project_code, data_date, predicted_finish_date, expected_delay_days, p_finish_after_contract, drivers
    from f
    order by expected_delay_days desc nulls last, project_code
    limit %s;
    """
    insert_sql = base_sql + """
    insert into analytics.forecasts (
      project_id, forecast_date, forecast_type, model_name, model_version, horizon_days,
      predicted_finish_date, p_finish_after_contract, expected_delay_days, p50_delay_days, p80_delay_days, p95_delay_days,
      drivers, input_feature_snapshot
    )
    select
      f.project_id,
      current_date,
      'project_finish',
      'heuristic_schedule_delay',
      'v0.1',
      180,
      f.predicted_finish_date,
      f.p_finish_after_contract,
      f.expected_delay_days,
      f.expected_delay_days,
      round((f.expected_delay_days * 1.25)::numeric, 2),
      round((f.expected_delay_days * 1.60)::numeric, 2),
      f.drivers,
      f.feature_snapshot
    from f
    where not exists (
      select 1 from analytics.forecasts af
      where af.project_id = f.project_id
        and af.forecast_date = current_date
        and af.forecast_type = 'project_finish'
        and af.model_name = 'heuristic_schedule_delay'
        and af.model_version = 'v0.1'
    );
    """
    with db_conn() as conn, conn.cursor() as cur:
        inserted = None
        if args.execute:
            cur.execute(insert_sql)
            inserted = cur.rowcount
            conn.commit()
        cur.execute(preview_sql, (args.limit,))
        rows = cur.fetchall()
        if args.json:
            payload = {
                "mode": "execute" if args.execute else "preview",
                "inserted": inserted,
                "rows": [dict(zip(["project", "data_date", "predicted_finish", "expected_delay_days", "p_finish_after_contract", "drivers"], r)) for r in rows]
            }
            print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))
        else:
            if inserted is not None:
                print(f"inserted={inserted}")
            print_rows(rows, ["project", "data_date", "predicted_finish", "expected_delay_days", "p_finish_after_contract", "drivers"], False)


def cmd_executive_risk_report(args):
    sql = """
    with latest_forecasts as (
      select distinct on (project_id)
        project_id,
        forecast_date,
        predicted_finish_date,
        expected_delay_days,
        p_finish_after_contract,
        p80_delay_days,
        p95_delay_days
      from analytics.forecasts
      where forecast_type='project_finish'
        and model_name='heuristic_schedule_delay'
        and model_version='v0.1'
      order by project_id, forecast_date desc, created_at desc
    ), universe as (
      select m.*, p.contract_finish_date,
        case when m.project_code ~* '(^|_)TESTE($|_)|OBRA_MODELO|MODELO|TEMPLATE|SCHDULE|ARQUIVOS_AUXILIARES' then true else false end is_model_or_auxiliary
      from analytics.v_project_schedule_metrics_current m
      join schedule.projects p on p.id = m.project_id
    ), risk as (
      select
        u.project_id,
        u.project_code,
        u.project_name,
        u.version_label,
        u.data_date,
        u.activities_total,
        u.critical_total,
        u.critical_open,
        u.critical_overdue,
        u.activities_overdue,
        round(u.avg_percent_complete::numeric, 2) avg_percent_complete,
        u.max_current_finish,
        u.max_baseline_finish,
        u.projected_delay_days_vs_baseline,
        lf.predicted_finish_date,
        lf.expected_delay_days,
        lf.p_finish_after_contract,
        lf.p80_delay_days,
        lf.p95_delay_days,
        case
          when u.is_model_or_auxiliary then 'EXCLUIR'
          when coalesce(lf.expected_delay_days,0) >= 60 or coalesce(u.critical_overdue,0) > 0 then 'CRITICO'
          when coalesce(lf.expected_delay_days,0) >= 30 or coalesce(u.critical_open,0) >= 100 then 'ALTO'
          when coalesce(lf.expected_delay_days,0) >= 15 or coalesce(u.critical_open,0) >= 50 then 'MEDIO'
          else 'BAIXO'
        end as executive_risk,
        (
          greatest(coalesce(lf.expected_delay_days,0),0) * 2
          + coalesce(u.critical_overdue,0) * 10
          + coalesce(u.critical_open,0) * 0.5
          + coalesce(u.activities_overdue,0) * 2
        )::numeric(12,2) executive_score
      from universe u
      left join latest_forecasts lf on lf.project_id = u.project_id
      where not u.is_model_or_auxiliary
        and u.max_current_finish >= current_date
        and exists (select 1 from portfolio.project_registry pr where pr.company='Mensura' and pr.schedule_project_id=u.project_id and pr.include_in_executive_report is true)
    )
    select
      project_code,
      version_label,
      data_date,
      activities_total,
      critical_total,
      critical_open,
      critical_overdue,
      avg_percent_complete,
      max_baseline_finish,
      max_current_finish,
      projected_delay_days_vs_baseline,
      predicted_finish_date,
      expected_delay_days,
      p_finish_after_contract,
      p80_delay_days,
      p95_delay_days,
      executive_risk,
      executive_score,
      case
        when executive_risk='CRITICO' then 'Revisão executiva imediata: validar baseline, caminho crítico e plano de recuperação.'
        when executive_risk='ALTO' then 'Abrir análise de frente/WBS e confirmar mitigação de críticas abertas.'
        when executive_risk='MEDIO' then 'Monitorar semanalmente e exigir explicação para variação de término.'
        else 'Acompanhar em rotina normal.'
      end as recommended_action
    from risk
    order by executive_score desc nulls last, expected_delay_days desc nulls last, critical_open desc nulls last
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.limit,))
        rows = cur.fetchall()
    headers = ["project", "version", "data_date", "activities", "critical", "critical_open", "critical_overdue", "avg_pct", "baseline_finish", "current_finish", "delay_vs_baseline", "predicted_finish", "expected_delay_days", "p_finish_after_contract", "p80_delay", "p95_delay", "risk", "score", "recommended_action"]
    if args.json:
        print(json.dumps([dict(zip(headers, r)) for r in rows], ensure_ascii=False, indent=2, default=str))
        return
    print("# Mensura Schedule Control — Relatório Executivo de Risco")
    print("")
    print(f"Projetos no relatório: {len(rows)}")
    print("")
    for i, r in enumerate(rows, 1):
        d = dict(zip(headers, r))
        print(f"## {i}. {d['project']} — {d['risk']}")
        print(f"- Versão/data-base: {d['version']} / {d['data_date']}")
        print(f"- Atividades: {d['activities']} | Críticas: {d['critical']} | Críticas abertas: {d['critical_open']} | Críticas atrasadas: {d['critical_overdue']}")
        print(f"- Avanço médio: {d['avg_pct']}%")
        print(f"- Baseline finish: {d['baseline_finish']} | Current finish: {d['current_finish']} | Delay vs baseline: {d['delay_vs_baseline']} dias")
        print(f"- Forecast finish: {d['predicted_finish']} | Atraso esperado: {d['expected_delay_days']} dias | P(atraso): {d['p_finish_after_contract']}")
        print(f"- P80/P95 delay: {d['p80_delay']} / {d['p95_delay']} dias")
        print(f"- Ação: {d['recommended_action']}")
        print("")


def cmd_portfolio_seed(args):
    rows = [
        ('Mensura','P_G','P&G Louveira','ongoing',True,True,'schedule_control','P_G',None,'Projeto sensível: notificação legal ativa; comunicação externa exige dupla revisão do Alê.'),
        ('Mensura','MELICITA_R1','Melicita','ongoing',True,True,'schedule_control','MELICITA_R1',None,'Referência válida para Melicita no relatório executivo.'),
        ('Mensura','MELICITA','Melicita duplicado','duplicate',False,False,'schedule_control','MELICITA','MELICITA_R1','Duplicado de MELICITA_R1 por decisão do Alê.'),
        ('Mensura','DOPPIO','Doppio','excluded',False,False,'schedule_control','DOPPIO',None,'Excluído do relatório executivo por decisão operacional do Alê.'),
        ('Mensura','ELEV_ALTO_DO_IPIRANGA','ELEV Alto do Ipiranga','excluded',False,False,'schedule_control','ELEV_ALTO_DO_IPIRANGA',None,'Excluído do relatório executivo por decisão operacional do Alê.'),
        ('Mensura','SOFITEL_DIRETOR','Sofitel Diretor','excluded',False,False,'schedule_control','SOFITEL_DIRETOR',None,'Excluído do relatório executivo por decisão operacional do Alê.'),
        ('Mensura','CCN_BIOMA','CCN Bioma','excluded',False,False,'schedule_control','CCN_BIOMA',None,'Excluído do relatório executivo por decisão operacional do Alê.'),
        ('Mensura','DF345_DIOGO_DE_FARIA','DF345 Diogo de Faria','excluded',False,False,'schedule_control','DF345_DIOGO_DE_FARIA',None,'Excluído do relatório executivo por decisão operacional do Alê.'),
        ('MIA','CCSP_CASA_7','CCSP Casa 7','ongoing',True,False,'2nd_brain',None,None,'Obra MIA ativa; ainda fora do Supabase Schedule Control.'),
        ('PCS','TEATRO_SUZANO','Teatro Município de Suzano','candidate',False,False,'sienge',None,None,'Candidato PCS; validar se obra ativa e cronograma.'),
        ('PCS','PARANAPIACABA','Paranapiacaba / Pavimentação Paranapiacaba','candidate',False,False,'sienge',None,None,'Candidato PCS; patrimônio sensível/CONDEPHAAT, validar status ativo.'),
        ('PCS','SPTRANS_1_OS','SPTrans — 1ª OS','candidate',False,False,'2nd_brain',None,None,'Aguardando recebimento da OS.'),
    ]
    sql = """
    insert into portfolio.project_registry (
      company, canonical_code, canonical_name, operational_status,
      include_in_executive_report, include_in_forecast, source_system,
      schedule_project_id, canonical_project_id, exclusion_reason, notes, metadata
    )
    values (
      %s,%s,%s,%s,%s,%s,%s,
      (select id from schedule.projects where code=%s),
      (select id from portfolio.project_registry where company=%s and canonical_code=%s),
      case when %s in ('excluded','duplicate') then %s else null end,
      %s,
      jsonb_build_object('seeded_by','mensura-schedule portfolio-seed','seeded_at',now())
    )
    on conflict (company, canonical_code) do update set
      canonical_name=excluded.canonical_name,
      operational_status=excluded.operational_status,
      include_in_executive_report=excluded.include_in_executive_report,
      include_in_forecast=excluded.include_in_forecast,
      source_system=excluded.source_system,
      schedule_project_id=excluded.schedule_project_id,
      canonical_project_id=excluded.canonical_project_id,
      exclusion_reason=excluded.exclusion_reason,
      notes=excluded.notes,
      metadata=portfolio.project_registry.metadata || excluded.metadata;
    """
    alias_sql = """
    insert into portfolio.project_aliases(project_registry_id, alias, alias_type, source_system)
    select id, %s, %s, %s from portfolio.project_registry where company=%s and canonical_code=%s
    on conflict do nothing;
    """
    with db_conn() as conn, conn.cursor() as cur:
        for company, code, name, status, inc_exec, inc_fc, source, sched_code, parent, note in rows:
            cur.execute(sql, (company, code, name, status, inc_exec, inc_fc, source, sched_code, company, parent or code, status, note, note))
            cur.execute(alias_sql, (code, 'code', source, company, code))
            cur.execute(alias_sql, (name, 'name', source, company, code))
            if code == 'P_G':
                cur.execute(alias_sql, ('P&G Louveira', 'name', 'manual', company, code))
            if code == 'CCSP_CASA_7':
                cur.execute(alias_sql, ('Casa 7', 'name', 'manual', company, code))
        conn.commit()
    payload={'mode':'execute','upserted':len(rows)}
    print(json.dumps(payload,ensure_ascii=False,indent=2) if args.json else payload)


def cmd_portfolio_registry(args):
    sql = """
    select company, canonical_code, canonical_name, operational_status,
      include_in_executive_report, include_in_forecast, source_system,
      schedule_project_code, canonical_parent_code, exclusion_reason, notes, aliases
    from portfolio.v_project_registry_current
    where (%s is null or company=%s)
      and (%s is null or operational_status=%s)
    order by company, case operational_status when 'ongoing' then 1 when 'candidate' then 2 when 'duplicate' then 3 when 'excluded' then 4 else 5 end, canonical_code
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.company,args.company,args.status,args.status,args.limit))
        print_rows(cur.fetchall(), ['company','code','name','status','executive','forecast','source','schedule_code','parent','exclusion_reason','notes','aliases'], args.json)



def _mpp_jdate(x):
    return None if x is None else str(x)[:10]


def _mpp_duration_days(x):
    if x is None:
        return None
    try:
        return float(x.getDuration())
    except Exception:
        return None


def _mpp_num_value(x):
    if x is None:
        return None
    try:
        return float(x)
    except Exception:
        return str(x)


def _mpp_status(percent):
    if percent == 100:
        return "completed"
    if percent == 0:
        return "not_started"
    return "in_progress"


def _read_mpp_project(path: Path):
    try:
        import mpxj
        from jpype import JClass
    except Exception as exc:
        raise SystemExit(f"MPXJ/Jpype indisponível para leitura .mpp: {exc}")
    if not mpxj.isJVMStarted():
        mpxj.startJVM("-Dlog4j2.StatusLogger.level=OFF", "-Dorg.apache.logging.log4j.simplelog.StatusLogger.level=OFF")
    return JClass("org.mpxj.reader.UniversalProjectReader")().read(str(path))


def _valid_mpp_tasks(project):
    return [t for t in list(project.getTasks()) if t is not None and t.getID() is not None and int(t.getID()) != 0 and t.getName()]



def cmd_intake_packages(args):
    sql = """
    select package_code, company, package_status, lot_number, expected_filename,
      intake_status, received_size_mb, item_count, imported_count, failed_count, partial_count, notes
    from intake.v_schedule_intake_dashboard
    where (%s is null or company=%s)
      and (%s is null or package_status=%s)
    order by company, package_status, lot_number
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.company, args.company, args.status, args.status, args.limit))
        print_rows(cur.fetchall(), ['package_code','company','package_status','lot','expected_filename','intake_status','received_size_mb','item_count','imported_count','failed_count','partial_count','notes'], args.json)


def cmd_intake_package_items(args):
    sql = """
    select p.package_code, i.company, i.project_code, i.project_name,
      i.project_operational_status, i.relative_path, i.source_filename, i.file_type,
      i.version_label, i.data_date, i.item_status, i.is_partial, i.schedule_version_id, i.notes
    from intake.schedule_package_items i
    join intake.schedule_packages p on p.id=i.package_id
    where (%s is null or p.package_code=%s)
      and (%s is null or i.company=%s)
      and (%s is null or i.item_status=%s)
    order by p.package_code, i.project_code, i.relative_path
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.package_code, args.package_code, args.company, args.company, args.item_status, args.item_status, args.limit))
        print_rows(cur.fetchall(), ['package_code','company','project_code','project_name','project_operational_status','relative_path','source_filename','file_type','version_label','data_date','item_status','is_partial','schedule_version_id','notes'], args.json)

def cmd_import_mpp(args):
    file_path = Path(args.file).expanduser().resolve()
    if not file_path.exists():
        raise SystemExit(f"Arquivo não encontrado: {file_path}")
    if file_path.read_bytes()[:8].hex() != "d0cf11e0a1b11ae1":
        raise SystemExit("Arquivo não parece Microsoft Project .mpp/CFB válido")
    project = _read_mpp_project(file_path)
    props = project.getProjectProperties()
    tasks = _valid_mpp_tasks(project)
    if not tasks:
        raise SystemExit("Nenhuma tarefa válida encontrada no .mpp")
    file_hash = hashlib.sha256(file_path.read_bytes()).hexdigest()
    data_date = args.data_date or _mpp_jdate(props.getStatusDate()) or _mpp_jdate(props.getCurrentDate())
    start_date = _mpp_jdate(props.getStartDate())
    finish_date = _mpp_jdate(props.getFinishDate())
    version_label = args.version_label or f"mpp_{data_date}_{file_hash[:8]}"
    critical_count = sum(1 for t in tasks if t.getCritical() is not None and bool(t.getCritical()))
    pct_values = []
    for t in tasks:
        v = _mpp_num_value(t.getPercentageComplete())
        if isinstance(v, float):
            pct_values.append(v)
    avg_percent = round(sum(pct_values)/len(pct_values), 2) if pct_values else None
    summary = {
        "mode": "execute" if args.execute else "preview",
        "file": str(file_path),
        "file_hash": file_hash[:16] + "...",
        "company": args.company,
        "project_code": args.project_code,
        "project_name": args.project_name,
        "version_label": version_label,
        "data_date": data_date,
        "start": start_date,
        "finish": finish_date,
        "tasks": len(tasks),
        "critical_tasks": critical_count,
        "avg_percent": avg_percent,
        "original_title": str(props.getProjectTitle()) if props.getProjectTitle() else None,
        "warnings": [],
    }
    if not data_date:
        summary["warnings"].append("data_date ausente; informe --data-date")
    if not start_date or not finish_date:
        summary["warnings"].append("start/finish do projeto ausente no arquivo")
    if critical_count == 0:
        summary["warnings"].append("nenhuma atividade crítica encontrada")
    if not args.execute:
        print(json.dumps(summary, ensure_ascii=False, indent=2) if args.json else summary)
        return
    if summary["warnings"] and not args.allow_warnings:
        raise SystemExit("Warnings encontrados; revise preview ou use --allow-warnings")

    with db_conn() as conn, conn.cursor() as cur:
        cur.execute("""
            insert into schedule.projects(code,name,company,client_name,contract_start_date,contract_finish_date,status,metadata)
            values (%s,%s,%s,%s,%s,%s,'active',jsonb_build_object('source','import_mpp','original_title',%s,'file_hash',%s))
            on conflict (code) do update set
              name=excluded.name,
              company=excluded.company,
              client_name=coalesce(excluded.client_name, schedule.projects.client_name),
              contract_start_date=coalesce(excluded.contract_start_date, schedule.projects.contract_start_date),
              contract_finish_date=coalesce(excluded.contract_finish_date, schedule.projects.contract_finish_date),
              status='active',
              metadata=schedule.projects.metadata || excluded.metadata
            returning id;
        """, (args.project_code, args.project_name, args.company, args.client_name, start_date, finish_date, summary["original_title"], file_hash))
        project_id = cur.fetchone()[0]
        cur.execute("select id from schedule.schedule_versions where project_id=%s and version_label=%s", (project_id, version_label))
        if cur.fetchone() and not args.force_new:
            raise SystemExit(f"Versão já existe para este projeto: {version_label}. Use outro --version-label ou --force-new.")
        cur.execute("""
            insert into raw.import_jobs(project_id,source_tool,source_filename,source_uri,file_hash,imported_by,status,data_date,raw_metadata)
            values (%s,'ms_project',%s,%s,%s,'Flávia/OpenClaw','parsed',%s,jsonb_build_object('tasks',%s,'status_date',%s,'project_title',%s,'import_command','import-mpp'))
            returning id;
        """, (project_id, file_path.name, str(file_path), file_hash, data_date, len(tasks), data_date, summary["original_title"]))
        import_job_id = cur.fetchone()[0]
        cur.execute("""
            insert into schedule.schedule_versions(project_id,import_job_id,version_label,data_date,source_tool,is_baseline_candidate,schedule_quality_score,notes,metadata)
            values (%s,%s,%s,%s,'ms_project',%s,%s,%s,jsonb_build_object('file_hash',%s,'original_path',%s,'import_command','import-mpp'))
            returning id;
        """, (project_id, import_job_id, version_label, data_date, args.baseline_candidate, args.quality_score, args.notes, file_hash, str(file_path)))
        version_id = cur.fetchone()[0]
        inserted = 0
        for idx, t in enumerate(tasks, start=1):
            uid = str(t.getUniqueID()) if t.getUniqueID() is not None else f"id-{t.getID()}"
            name = str(t.getName())
            activity_code = str(t.getOutlineNumber() or t.getID())
            wbs = str(t.getWBS() or t.getOutlineNumber() or "") or None
            pct = _mpp_num_value(t.getPercentageComplete())
            raw = {
                "id": int(t.getID()) if t.getID() is not None else None,
                "unique_id": uid,
                "name": name,
                "outline_number": str(t.getOutlineNumber() or ""),
                "outline_level": int(t.getOutlineLevel()) if t.getOutlineLevel() is not None else None,
                "wbs": str(t.getWBS() or ""),
                "start": str(t.getStart()) if t.getStart() else None,
                "finish": str(t.getFinish()) if t.getFinish() else None,
                "baseline_start": str(t.getBaselineStart()) if t.getBaselineStart() else None,
                "baseline_finish": str(t.getBaselineFinish()) if t.getBaselineFinish() else None,
                "actual_start": str(t.getActualStart()) if t.getActualStart() else None,
                "actual_finish": str(t.getActualFinish()) if t.getActualFinish() else None,
                "duration": str(t.getDuration()) if t.getDuration() else None,
                "remaining_duration": str(t.getRemainingDuration()) if t.getRemainingDuration() else None,
                "percent_complete": pct,
                "physical_percent_complete": _mpp_num_value(t.getPhysicalPercentComplete()),
                "critical": bool(t.getCritical()) if t.getCritical() is not None else None,
                "total_slack": str(t.getTotalSlack()) if t.getTotalSlack() else None,
                "free_slack": str(t.getFreeSlack()) if t.getFreeSlack() else None,
                "constraint_type": str(t.getConstraintType()) if t.getConstraintType() else None,
            }
            cur.execute("""
                insert into raw.source_rows(import_job_id,source_row_number,source_activity_uid,activity_code,raw_payload)
                values (%s,%s,%s,%s,%s::jsonb);
            """, (import_job_id, idx, uid, activity_code, json.dumps(raw, ensure_ascii=False)))
            cur.execute("""
                insert into schedule.activity_identities(project_id,source_activity_uid,activity_code,normalized_name,wbs_code,identity_confidence,first_seen_version_id)
                values (%s,%s,%s,%s,%s,'high',%s)
                on conflict (project_id, source_activity_uid) do update set
                  activity_code=excluded.activity_code,
                  normalized_name=excluded.normalized_name,
                  wbs_code=excluded.wbs_code,
                  is_active=true
                returning id;
            """, (project_id, uid, activity_code, norm(name), wbs, version_id))
            identity_id = cur.fetchone()[0]
            cur.execute("""
                insert into schedule.activity_versions(
                  schedule_version_id, activity_identity_id, activity_code, activity_name, activity_type, wbs_code,
                  planned_duration_days, remaining_duration_days, baseline_start, baseline_finish,
                  current_start, current_finish, actual_start, actual_finish, percent_complete, physical_percent_complete,
                  is_critical, total_float_days, free_float_days, constraint_type, status, raw_payload
                ) values (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s::jsonb);
            """, (version_id, identity_id, activity_code, name, str(t.getType()) if t.getType() else None, wbs,
                  _mpp_duration_days(t.getDuration()), _mpp_duration_days(t.getRemainingDuration()), _mpp_jdate(t.getBaselineStart()), _mpp_jdate(t.getBaselineFinish()),
                  _mpp_jdate(t.getStart()), _mpp_jdate(t.getFinish()), _mpp_jdate(t.getActualStart()), _mpp_jdate(t.getActualFinish()), pct, _mpp_num_value(t.getPhysicalPercentComplete()),
                  bool(t.getCritical()) if t.getCritical() is not None else None, _mpp_duration_days(t.getTotalSlack()), _mpp_duration_days(t.getFreeSlack()), str(t.getConstraintType()) if t.getConstraintType() else None,
                  _mpp_status(pct), json.dumps(raw, ensure_ascii=False)))
            inserted += 1
        cur.execute("""
            insert into portfolio.project_registry(company, canonical_code, canonical_name, operational_status, include_in_executive_report, include_in_forecast, source_system, schedule_project_id, notes, metadata)
            values (%s,%s,%s,'ongoing',%s,%s,'mixed',%s,%s,jsonb_build_object('last_import_version_label',%s,'last_import_file_hash',%s))
            on conflict (company, canonical_code) do update set
              canonical_name=excluded.canonical_name,
              operational_status='ongoing',
              include_in_executive_report=excluded.include_in_executive_report,
              include_in_forecast=excluded.include_in_forecast,
              source_system='mixed',
              schedule_project_id=excluded.schedule_project_id,
              notes=coalesce(excluded.notes, portfolio.project_registry.notes),
              metadata=portfolio.project_registry.metadata || excluded.metadata;
        """, (args.company, args.project_code, args.project_name, args.include_executive, args.include_forecast, project_id, args.registry_notes, version_label, file_hash))
        conn.commit()
    summary.update({"project_id": str(project_id), "version_id": str(version_id), "import_job_id": str(import_job_id), "tasks_imported": inserted})
    print(json.dumps(summary, ensure_ascii=False, indent=2) if args.json else summary)

def cmd_portfolio_classify_pcs(args):
    # Classificação estrutural inicial da base PCS/Sienge informada pelo Alê.
    # Como as obras Sienge não têm cronograma, nenhuma entra em forecast/executive até haver cronograma importado.
    records = [
        ('14','FACULDADE DE DIREITO DE SBC - CENTRO JURIDICO','candidate','Registro Sienge; validar se obra ativa, encerrada ou apenas referência histórica.'),
        ('16','RETROFIT FACULDADE DE DIREITO DE SBC','candidate','Registro Sienge; validar status operacional.'),
        ('24','MANUTENÇÃO PARANAPIACABA 2024','candidate','Registro Sienge; validar se é frente encerrada ou ativa.'),
        ('71','CUSTOS COM PESSOAL - DIRETORIA','administrative','Centro administrativo; não tratar como obra.'),
        ('83','SMARTJAMPA','candidate','Registro Sienge; validar status operacional.'),
        ('100','ADM PCS','administrative','Centro administrativo PCS; não tratar como obra.'),
        ('181','RUA LAURA 181','candidate','Registro Sienge; validar status operacional.'),
        ('1000','ADM MIA ENGENHARIA','administrative','Centro administrativo MIA; não tratar como obra PCS.'),
        ('1354','TEATRO MUNICÍPIO DE SUZANO','candidate','Obra PCS; cronograma em elaboração; fora de forecast/executive até importação.'),
        ('1900','REFORMA E MAN. DE ESCOLAS SP - LOTE 180','candidate','Registro Sienge; validar contrato/frente ativa.'),
        ('1901','EMEI EDUARDO CARLOS PEREIRA','candidate','Possível unidade do lote escolas; validar relação com contrato principal.'),
        ('1902','EMEI CIDADE FERNÃO DIAS','candidate','Possível unidade do lote escolas; validar relação com contrato principal.'),
        ('1903','EMEI AVIADOR EDU CHAVES','candidate','Possível unidade do lote escolas; validar relação com contrato principal.'),
        ('1904','EMEF PROFA. ESMERALDA SALLES PEREIRA','candidate','Possível unidade do lote escolas; validar relação com contrato principal.'),
        ('1905','EMEI PROF. ÊNIO CORREA','candidate','Possível unidade do lote escolas; validar relação com contrato principal.'),
        ('4412','GERENCIAL PAVIMENTAÇÃO PARANAPICABA PCS','candidate','Registro gerencial; validar grafia e relação com Paranapiacaba.'),
        ('4428','AV. NOVE DE JULHO 4428','candidate','Registro Sienge; validar status operacional.'),
        ('4500','ATA CALÇADAS','candidate','Registro Sienge; validar se contrato/frente ativa.'),
        ('5201','MP AUDITÓRIO RIACHUELO','candidate','Registro Sienge; validar status operacional.'),
        ('45000','ATA CALÇADAS DIARIO DE OBRA 01 A 12 DE MARÇO','candidate','Provável registro operacional/diário; validar se não é duplicado de ATA CALÇADAS.'),
        ('94539','LOCAÇÃO SÃO CAETANO - VUC PCS','administrative','Locação/equipamento; não tratar como obra física.'),
        ('441','PAVIMENTAÇÃO PARANAPIACABA','candidate','Obra/frente PCS; validar status e cronograma futuro.'),
    ]
    if not args.execute:
        rows=[]
        for code,name,status,note in records:
            rows.append((code,name,status,False,False,note))
        print_rows(rows, ['sienge_code','name','suggested_status','executive','forecast','note'], args.json)
        return
    sql = """
    insert into portfolio.project_registry(company, canonical_code, canonical_name, operational_status, include_in_executive_report, include_in_forecast, source_system, notes, metadata)
    values ('PCS', %s, %s, %s, false, false, 'sienge', %s, jsonb_build_object('sienge_code', %s, 'schedule_status', 'not_available', 'classified_by', 'mensura-schedule portfolio-classify-pcs'))
    on conflict (company, canonical_code) do update set
      canonical_name=excluded.canonical_name,
      operational_status=excluded.operational_status,
      include_in_executive_report=false,
      include_in_forecast=false,
      source_system='sienge',
      notes=excluded.notes,
      metadata=portfolio.project_registry.metadata || excluded.metadata;
    """
    alias_sql = """
    insert into portfolio.project_aliases(project_registry_id, alias, alias_type, source_system)
    select id, %s, %s, 'sienge' from portfolio.project_registry where company='PCS' and canonical_code=%s
    on conflict do nothing;
    """
    with db_conn() as conn, conn.cursor() as cur:
        for code,name,status,note in records:
            canonical_code = 'PCS_SIENGE_' + code
            if code == '1354': canonical_code = 'TEATRO_SUZANO'
            elif code in ('441','4412'): canonical_code = 'PARANAPIACABA' if code == '441' else 'PARANAPIACABA_GERENCIAL'
            cur.execute(sql, (canonical_code, name, status, note, code))
            cur.execute(alias_sql, (code, 'sienge', canonical_code))
            cur.execute(alias_sql, (name, 'name', canonical_code))
        conn.commit()
    payload={'mode':'execute','classified':len(records),'executive_included':0,'forecast_included':0}
    print(json.dumps(payload, ensure_ascii=False, indent=2) if args.json else payload)



def cmd_portfolio_risk_summary(args):
    sql = """
    with executive_projects as (
      select pr.company, pr.canonical_code, pr.canonical_name, pr.schedule_project_id
      from portfolio.project_registry pr
      where pr.include_in_executive_report is true
        and pr.schedule_project_id is not null
        and (%s is null or pr.company=%s)
    ), latest_versions as (
      select distinct on (sv.project_id)
        sv.project_id, sv.id as schedule_version_id, sv.version_label, sv.data_date
      from schedule.schedule_versions sv
      join executive_projects ep on ep.schedule_project_id=sv.project_id
      order by sv.project_id, sv.data_date desc, sv.created_at desc
    ), metrics as (
      select
        ep.company,
        ep.canonical_code,
        ep.canonical_name,
        lv.version_label,
        lv.data_date,
        count(av.*)::int activities_total,
        count(*) filter (where av.is_critical)::int critical_total,
        count(*) filter (where av.is_critical and coalesce(av.percent_complete,0) < 100)::int critical_open,
        count(*) filter (where av.is_critical and coalesce(av.percent_complete,0) < 100 and av.current_finish < lv.data_date)::int critical_overdue,
        count(*) filter (where coalesce(av.percent_complete,0) < 100 and av.current_finish < lv.data_date)::int activities_overdue,
        round(avg(av.percent_complete)::numeric,2) real_percent,
        round(avg(
          case
            when av.baseline_start is null or av.baseline_finish is null then null
            when lv.data_date < av.baseline_start then 0
            when lv.data_date >= av.baseline_finish then 100
            when av.baseline_finish = av.baseline_start then 100
            else greatest(0, least(100, (100.0 * (lv.data_date - av.baseline_start)::numeric / nullif((av.baseline_finish - av.baseline_start),0))))
          end
        )::numeric,2) planned_percent,
        min(av.current_start) current_start,
        max(av.current_finish) current_finish,
        min(av.baseline_start) baseline_start,
        max(av.baseline_finish) baseline_finish
      from executive_projects ep
      join latest_versions lv on lv.project_id=ep.schedule_project_id
      join schedule.activity_versions av on av.schedule_version_id=lv.schedule_version_id
      group by ep.company, ep.canonical_code, ep.canonical_name, lv.version_label, lv.data_date
    ), latest_forecasts as (
      select distinct on (project_id)
        project_id, expected_delay_days, predicted_finish_date, p_finish_after_contract
      from analytics.forecasts
      where forecast_type='project_finish'
        and model_name='heuristic_schedule_delay'
        and model_version='v0.1'
      order by project_id, forecast_date desc, created_at desc
    ), scored as (
      select
        m.*,
        (m.current_finish - m.baseline_finish)::int delay_days,
        (m.real_percent - m.planned_percent)::numeric(12,2) percent_gap,
        lf.predicted_finish_date,
        coalesce(lf.expected_delay_days, greatest((m.current_finish - m.baseline_finish),0))::numeric expected_delay_days,
        lf.p_finish_after_contract,
        case
          when coalesce(lf.expected_delay_days, greatest((m.current_finish - m.baseline_finish),0)) >= 60 or m.critical_overdue > 0 or (m.real_percent - m.planned_percent) <= -20 then 'CRITICO'
          when coalesce(lf.expected_delay_days, greatest((m.current_finish - m.baseline_finish),0)) >= 30 or m.critical_open >= 50 or (m.real_percent - m.planned_percent) <= -10 then 'ALTO'
          when coalesce(lf.expected_delay_days, greatest((m.current_finish - m.baseline_finish),0)) >= 15 or m.critical_open >= 20 or (m.real_percent - m.planned_percent) <= -5 then 'MEDIO'
          else 'BAIXO'
        end as executive_risk
      from metrics m
      left join schedule.projects sp on sp.code=m.canonical_code
      left join latest_forecasts lf on lf.project_id=sp.id
    )
    select
      company,
      canonical_code,
      canonical_name,
      version_label,
      data_date,
      baseline_finish,
      current_finish,
      delay_days,
      planned_percent,
      real_percent,
      percent_gap,
      activities_total,
      critical_total,
      critical_open,
      critical_overdue,
      activities_overdue,
      predicted_finish_date,
      expected_delay_days,
      p_finish_after_contract,
      executive_risk,
      case
        when executive_risk='CRITICO' then 'Diretoria: exigir plano de recuperação, dono claro e decisão sobre caminho crítico em até 48h.'
        when executive_risk='ALTO' then 'Diretoria: revisar frentes críticas e cobrar plano semanal de mitigação.'
        when executive_risk='MEDIO' then 'Gestão: monitorar semanalmente e antecipar restrições antes de virar atraso executivo.'
        else 'Manter rotina de acompanhamento; sem intervenção executiva imediata.'
      end as recommendation
    from scored
    order by case executive_risk when 'CRITICO' then 1 when 'ALTO' then 2 when 'MEDIO' then 3 else 4 end,
      expected_delay_days desc nulls last,
      critical_open desc nulls last,
      company,
      canonical_code
    limit %s;
    """
    headers = ['company','project','name','version','data_date','baseline_finish','current_finish','delay_days','planned_percent','real_percent','percent_gap','activities','critical_total','critical_open','critical_overdue','activities_overdue','predicted_finish','expected_delay_days','p_finish_after_contract','risk','recommendation']
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.company, args.company, args.limit))
        rows = cur.fetchall()
    if args.json:
        print(json.dumps([dict(zip(headers, r)) for r in rows], ensure_ascii=False, indent=2, default=str))
        return
    print('# Relatório de Diretoria — Portfólio de Obras')
    print('')
    print(f'Obras no relatório: {len(rows)}')
    print('')
    for i, row in enumerate(rows, 1):
        d = dict(zip(headers, row))
        delay = d['delay_days'] if d['delay_days'] is not None else 'n/d'
        gap = d['percent_gap'] if d['percent_gap'] is not None else 'n/d'
        print(f"## {i}. {d['company']} — {d['project']} — {d['risk']}")
        print(f"- Baseline x término atual: {d['baseline_finish']} → {d['current_finish']} ({delay} dias)")
        print(f"- % planejado x % real: {d['planned_percent']}% → {d['real_percent']}% (gap {gap} p.p.)")
        print(f"- Caminho crítico: {d['critical_open']} críticas abertas; {d['critical_overdue']} críticas atrasadas")
        print(f"- Atividades: {d['activities']} totais; {d['activities_overdue']} atrasadas")
        if d['predicted_finish']:
            print(f"- Forecast: término previsto {d['predicted_finish']} | atraso esperado {d['expected_delay_days']} dias | P(atraso) {d['p_finish_after_contract']}")
        print(f"- Recomendação: {d['recommendation']}")
        print('')

def cmd_portfolio_report(args):
    sql = """
    select
      pr.company,
      pr.canonical_code,
      pr.canonical_name,
      pr.operational_status,
      pr.include_in_executive_report,
      pr.include_in_forecast,
      pr.source_system,
      sp.code as schedule_code,
      latest.version_label,
      latest.data_date,
      latest.activities,
      latest.critical,
      latest.max_finish,
      latest.avg_percent,
      pr.notes
    from portfolio.project_registry pr
    left join schedule.projects sp on sp.id = pr.schedule_project_id
    left join lateral (
      select sv.version_label, sv.data_date, count(av.*)::int activities,
        count(*) filter (where av.is_critical)::int critical,
        max(av.current_finish) max_finish,
        round(avg(av.percent_complete)::numeric, 2) avg_percent
      from schedule.schedule_versions sv
      left join schedule.activity_versions av on av.schedule_version_id=sv.id
      where sv.project_id=sp.id
      group by sv.id, sv.version_label, sv.data_date
      order by sv.created_at desc
      limit 1
    ) latest on true
    where (%s is null or pr.company=%s)
      and (%s is false or pr.include_in_executive_report is true)
    order by pr.company,
      case pr.operational_status when 'ongoing' then 1 when 'candidate' then 2 when 'administrative' then 3 when 'duplicate' then 4 when 'excluded' then 5 else 9 end,
      pr.canonical_code
    limit %s;
    """
    with db_conn() as conn, conn.cursor() as cur:
        cur.execute(sql, (args.company, args.company, args.executive_only, args.limit))
        print_rows(cur.fetchall(), ['company','code','name','status','executive','forecast','source','schedule_code','version','data_date','activities','critical','max_finish','avg_percent','notes'], args.json)

def cmd_validate(args):
    subprocess.run([sys.executable, str(ROOT / "scripts/validate_migration.py")], check=True)


def cmd_migration_list(args):
    load_env_file()
    env = os.environ.copy()
    subprocess.run(["/root/.local/bin/supabase", "migration", "list"], cwd=ROOT, env=env, check=True)


def cmd_lint(args):
    load_env_file()
    env = os.environ.copy()
    cmd = ["/root/.local/bin/supabase", "db", "lint", "--linked", "--schema", "raw,schedule,control,analytics"]
    subprocess.run(cmd, cwd=ROOT, env=env, check=True)


def cmd_ingest_sharepoint(args):
    load_env_file()
    env = os.environ.copy()
    cmd = [sys.executable, str(ROOT / "scripts/ingest_sharepoint_schedules.py")]
    if args.dry_run:
        cmd.append("--dry-run")
    if args.limit:
        cmd += ["--limit", str(args.limit)]
    subprocess.run(cmd, cwd=ROOT, env=env, check=True)


def graph_token(config_path: Path = MENSURA_GRAPH_CONFIG) -> str:
    cfg = json.loads(config_path.read_text(encoding="utf-8"))
    r = requests.post(
        f"https://login.microsoftonline.com/{cfg['tenantId']}/oauth2/v2.0/token",
        data={
            "client_id": cfg["clientId"],
            "client_secret": cfg["client" + "Secret"],
            "scope": cfg.get("scope", "https://graph.microsoft.com/.default"),
            "grant_type": "client_credentials",
        },
        timeout=30,
    )
    r.raise_for_status()
    return r.json()["access_token"]


def cmd_search_sharepoint(args):
    token = graph_token()
    queries = args.query or [
        '"Arquivos Auxiliares" xlsx schedule',
        '"Arquivos Auxiliares" schedule',
        'filename:xlsx schedule',
        'schedule xlsx',
    ]
    results: list[dict[str, Any]] = []
    for q in queries:
        body = {
            "requests": [{
                "entityTypes": ["driveItem"],
                "query": {"queryString": q},
                "from": 0,
                "size": args.size,
                "region": "BRA",
                "fields": ["name", "webUrl", "fileExtension", "parentReference", "size", "lastModifiedDateTime"],
            }]
        }
        r = requests.post(f"{GRAPH}/search/query", headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"}, json=body, timeout=60)
        if r.status_code >= 400:
            print(f"WARN query failed {q}: {r.status_code} {r.text[:240]}", file=sys.stderr)
            continue
        for resp in r.json().get("value", []):
            for cont in resp.get("hitsContainers", []):
                for hit in cont.get("hits", []):
                    res = hit.get("resource", {})
                    name = res.get("name", "")
                    web = res.get("webUrl") or ""
                    if not name.lower().endswith((".xlsx", ".xlsm", ".xls")):
                        continue
                    if args.contains and args.contains.lower() not in (name + " " + web).lower():
                        continue
                    results.append({
                        "query": q,
                        "name": name,
                        "webUrl": web,
                        "size": res.get("size"),
                        "lastModified": res.get("lastModifiedDateTime"),
                    })
    seen = set(); dedup = []
    for x in results:
        k = x.get("webUrl") or x.get("name")
        if k in seen:
            continue
        seen.add(k); dedup.append(x)
    out = Path(args.output) if args.output else None
    if out:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(dedup, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(dedup, ensure_ascii=False, indent=2) if args.json else f"found={len(dedup)} output={out or '-'}")
    if not args.json:
        for i, x in enumerate(dedup[: args.print_limit], 1):
            print(f"{i}. {x['name']} | {x.get('size')} | {x.get('lastModified')}\n   {x.get('webUrl')}")


def cmd_inspect_workbooks(args):
    manifest = json.loads(Path(args.input).read_text(encoding="utf-8"))
    out = []
    for m in manifest:
        lp = m.get("local_path")
        if not lp or not Path(lp).exists():
            out.append({**m, "inspection_status": "missing_local_path", "tables": []})
            continue
        try:
            wb = load_workbook(lp, read_only=False, data_only=True)
            tables = []
            for ws in wb.worksheets:
                for table in ws.tables.values():
                    tables.append({"sheet": ws.title, "table": table.name, "ref": table.ref})
            out.append({**m, "inspection_status": "ok", "sheet_count": len(wb.worksheets), "tables": tables})
        except Exception as e:
            out.append({**m, "inspection_status": "error", "error": str(e)[:500], "tables": []})
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    schedule = [x for x in out if any("schedule" in (t.get("table", "").lower()) for t in x.get("tables", []))]
    print(f"inspected={len(out)} schedule_table_workbooks={len(schedule)} output={output}")
    if args.print_schedule:
        for i, x in enumerate(schedule[: args.print_limit], 1):
            ts = ", ".join(f"{t['table']}@{t['sheet']}:{t['ref']}" for t in x.get("tables", []) if "schedule" in t.get("table", "").lower())
            print(f"{i}. {x.get('name')} | {ts}")


def cmd_download_sharepoint(args):
    token = graph_token()
    items = json.loads(Path(args.input).read_text(encoding="utf-8"))
    outdir = Path(args.output_dir)
    outdir.mkdir(parents=True, exist_ok=True)
    manifest = []
    for i, item in enumerate(items, 1):
        url = item.get("webUrl")
        if not url:
            continue
        enc = base64.urlsafe_b64encode(url.encode()).decode().rstrip("=")
        meta = requests.get(f"{GRAPH}/shares/u!{enc}/driveItem", headers={"Authorization": f"Bearer {token}"}, timeout=60)
        if meta.status_code >= 400:
            manifest.append({**item, "download_status": meta.status_code, "error": meta.text[:300]})
            continue
        m = meta.json()
        drive_id = (m.get("parentReference") or {}).get("driveId")
        item_id = m.get("id")
        content = requests.get(f"{GRAPH}/drives/{drive_id}/items/{item_id}/content", headers={"Authorization": f"Bearer {token}"}, timeout=120, allow_redirects=True)
        safe = re.sub(r"[^A-Za-z0-9._ -]+", "_", item.get("name", f"file-{i}.xlsx")).strip()
        path = outdir / f"{i:02d}-{safe}"
        if content.status_code < 400:
            path.write_bytes(content.content)
            manifest.append({**item, "download_status": content.status_code, "local_path": str(path), "bytes": len(content.content), "driveId": drive_id, "graph_item_id": item_id})
        else:
            manifest.append({**item, "download_status": content.status_code, "error": content.text[:300], "driveId": drive_id, "graph_item_id": item_id})
    manifest_path = outdir.parent / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"downloaded={sum(1 for x in manifest if x.get('local_path'))} total={len(manifest)} manifest={manifest_path}")


def build_parser():
    p = argparse.ArgumentParser(prog="mensura-schedule", description="CLI interna do Mensura Schedule Control")
    sub = p.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("db-counts", help="Conta objetos principais no Supabase cloud")
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_db_counts)

    s = sub.add_parser("critical-summary", help="Resumo por projeto usando coluna Crítica do Project")
    s.add_argument("--limit", type=int, default=30)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_critical_summary)

    s = sub.add_parser("quality-report", help="Qualidade dos dados nas versões atuais")
    s.add_argument("--limit", type=int, default=50)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_quality_report)

    s = sub.add_parser("baseline-compare", help="Compara término atual x baseline_finish importado")
    s.add_argument("--project", help="Código do projeto")
    s.add_argument("--limit", type=int, default=50)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_baseline_compare)

    s = sub.add_parser("weekly-metrics-preview", help="Preview read-only de métricas semanais derivadas")
    s.add_argument("--limit", type=int, default=50)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_weekly_metrics_preview)

    s = sub.add_parser("audit-orphans", help="Audita raw.import_jobs sem schedule_version; não limpa")
    s.add_argument("--detail", action="store_true")
    s.add_argument("--limit", type=int, default=50)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_audit_orphans)


    s = sub.add_parser("project-classify", help="Classifica projetos como válidos, modelos, duplicados prováveis ou revisão")
    s.add_argument("--limit", type=int, default=80)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_project_classify)

    s = sub.add_parser("baseline-discovery", help="Descobre sinais de baseline por versão/projeto")
    s.add_argument("--only-candidates", action="store_true")
    s.add_argument("--limit", type=int, default=120)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_baseline_discovery)

    s = sub.add_parser("risk-report", help="Ranking executivo de risco dos cronogramas atuais")
    s.add_argument("--limit", type=int, default=50)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_risk_report)

    s = sub.add_parser("populate-weekly-metrics", help="Preview/upsert das métricas semanais atuais em analytics.project_weekly_metrics")
    s.add_argument("--execute", action="store_true")
    s.add_argument("--limit", type=int, default=50)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_populate_weekly_metrics)


    s = sub.add_parser("data-readiness", help="Semáforo de aptidão dos projetos para análise e predição")
    s.add_argument("--limit", type=int, default=80)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_data_readiness)


    s = sub.add_parser("repair-date-fields", help="Reparseia datas em raw_payload (ex.: 'Término' com dia da semana) e corrige campos nulos")
    s.add_argument("--execute", action="store_true")
    s.add_argument("--overwrite", action="store_true")
    s.add_argument("--limit", type=int, default=200000)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_repair_date_fields)


    s = sub.add_parser("analytics-universe", help="Lista universo analítico filtrando modelos/auxiliares")
    s.add_argument("--analytics-only", action="store_true")
    s.add_argument("--limit", type=int, default=80)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_analytics_universe)

    s = sub.add_parser("forecast-initial", help="Preview/insere forecast heurístico inicial de término por projeto")
    s.add_argument("--execute", action="store_true")
    s.add_argument("--limit", type=int, default=50)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_forecast_initial)


    s = sub.add_parser("executive-risk-report", help="Relatório executivo de risco filtrado por universo analítico e forecast")
    s.add_argument("--limit", type=int, default=20)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_executive_risk_report)


    s = sub.add_parser("intake-packages", help="Lista pacotes planejados/recebidos de cronogramas")
    s.add_argument("--company", choices=["Mensura", "MIA", "PCS", "Pessoal", "Outro"])
    s.add_argument("--status", choices=["em_andamento", "concluidas", "misto", "historico", "outro"])
    s.add_argument("--limit", type=int, default=100)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_intake_packages)

    s = sub.add_parser("intake-items", help="Lista itens de pacotes de cronogramas")
    s.add_argument("--package-code")
    s.add_argument("--company", choices=["Mensura", "MIA", "PCS", "Pessoal", "Outro"])
    s.add_argument("--item-status", choices=["planned", "received", "validated", "imported", "skipped", "failed", "partial"])
    s.add_argument("--limit", type=int, default=200)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_intake_package_items)

    s = sub.add_parser("import-mpp", help="Importa arquivo Microsoft Project .mpp como novo snapshot versionado")
    s.add_argument("file")
    s.add_argument("--company", required=True, choices=["Mensura", "MIA", "PCS", "Pessoal", "Outro"])
    s.add_argument("--project-code", required=True)
    s.add_argument("--project-name", required=True)
    s.add_argument("--client-name")
    s.add_argument("--version-label")
    s.add_argument("--data-date")
    s.add_argument("--baseline-candidate", action="store_true")
    s.add_argument("--quality-score", type=float, default=95)
    s.add_argument("--notes", default="Importação via mensura-schedule import-mpp")
    s.add_argument("--registry-notes")
    s.add_argument("--include-executive", action="store_true")
    s.add_argument("--include-forecast", action="store_true")
    s.add_argument("--force-new", action="store_true")
    s.add_argument("--allow-warnings", action="store_true")
    s.add_argument("--execute", action="store_true")
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_import_mpp)

    s = sub.add_parser("portfolio-classify-pcs", help="Classifica registros PCS/Sienge no portfolio registry")
    s.add_argument("--execute", action="store_true")
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_portfolio_classify_pcs)

    s = sub.add_parser("portfolio-risk-summary", help="Relatório sucinto de Diretoria: previsto x real, planejado x real, risco e recomendação")
    s.add_argument("--company", choices=["Mensura", "MIA", "PCS", "Pessoal", "Outro"])
    s.add_argument("--limit", type=int, default=50)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_portfolio_risk_summary)

    s = sub.add_parser("portfolio-report", help="Relatório transversal do portfolio registry com vínculo Schedule Control")
    s.add_argument("--company", choices=["Mensura", "MIA", "PCS", "Pessoal", "Outro"])
    s.add_argument("--executive-only", action="store_true")
    s.add_argument("--limit", type=int, default=200)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_portfolio_report)

    s = sub.add_parser("portfolio-seed", help="Aplica seed inicial do portfolio.project_registry")
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_portfolio_seed)

    s = sub.add_parser("portfolio-registry", help="Lista portfolio.project_registry")
    s.add_argument("--company", choices=["Mensura", "MIA", "PCS", "Pessoal", "Outro"])
    s.add_argument("--status")
    s.add_argument("--limit", type=int, default=100)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_portfolio_registry)

    s = sub.add_parser("validate", help="Valida estrutura da migration foundation")
    s.set_defaults(func=cmd_validate)

    s = sub.add_parser("migration-list", help="Lista migrations local/remoto")
    s.set_defaults(func=cmd_migration_list)

    s = sub.add_parser("lint", help="Roda lint remoto nos schemas principais")
    s.set_defaults(func=cmd_lint)

    s = sub.add_parser("ingest-sharepoint", help="Ingere workbooks baixados do runtime/sharepoint")
    s.add_argument("--dry-run", action="store_true")
    s.add_argument("--limit", type=int, default=0)
    s.set_defaults(func=cmd_ingest_sharepoint)

    s = sub.add_parser("search-sharepoint", help="Busca workbooks no SharePoint via Microsoft Graph Search")
    s.add_argument("--query", action="append", help="Query Graph Search; pode repetir")
    s.add_argument("--contains", default="", help="Filtro simples no nome/url")
    s.add_argument("--size", type=int, default=50)
    s.add_argument("--output", default=str(ROOT / "runtime/sharepoint/search_results.json"))
    s.add_argument("--print-limit", type=int, default=40)
    s.add_argument("--json", action="store_true")
    s.set_defaults(func=cmd_search_sharepoint)

    s = sub.add_parser("download-sharepoint", help="Baixa resultados de search-sharepoint")
    s.add_argument("--input", default=str(ROOT / "runtime/sharepoint/search_results.json"))
    s.add_argument("--output-dir", default=str(ROOT / "runtime/sharepoint/raw"))
    s.set_defaults(func=cmd_download_sharepoint)

    s = sub.add_parser("inspect-workbooks", help="Inspeciona abas/tabelas Excel baixadas")
    s.add_argument("--input", default=str(ROOT / "runtime/sharepoint/manifest.json"))
    s.add_argument("--output", default=str(ROOT / "runtime/sharepoint/workbook_inspection.json"))
    s.add_argument("--print-schedule", action="store_true")
    s.add_argument("--print-limit", type=int, default=80)
    s.set_defaults(func=cmd_inspect_workbooks)

    return p


def main(argv=None):
    args = build_parser().parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
